"""Shared plumbing every export provider uses.

**A file the deep-dive's §2 layout does not name**, added with a reason: seven providers all
need the same three things — read canonical rows for a user, build the workbook, store the
resulting artifact as a blob. Copying that into each provider is how seven formats drift into
seven subtly different views of the same data, which is the failure the framework exists to
prevent. The *format* logic stays in each provider; only the mechanism is here.

`openpyxl` is imported lazily and nowhere else in this package writes Excel. Its absence
degrades the workbook-producing providers to unavailable and leaves the CSV/IIF ones working
(`docs/PRINCIPLES.md` §3.3 point 5, §4.4).
"""

from __future__ import annotations

import io
import uuid
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import datetime
from typing import Any

from ...contracts import BlobRef, Receipt, StorageCodec, utcnow
from ...db.connection import Database
from ...db.receipts import ReceiptRepository
from ...historian.query import HistorianQuery
from ...blob_store.store import BlobStore
from ...reimport.parser import SNAPSHOT_CELL, SNAPSHOT_PREFIX, SNAPSHOT_SHEET
from .. import errors

#: The raw-data sheet's column order. One definition, used by the writer here and by
#: Reimport's parser through `receipt_id` — the two are a matched pair and a column added on
#: one side without the other is how a round trip silently loses a field.
EXPORT_COLUMNS = (
    "receipt_id",
    "transaction_date",
    "vendor_name",
    "currency",
    "total_amount",
    "vat_amount",
    "group_id",
)


@dataclass(frozen=True)
class ExportContext:
    """Everything a provider needs to read canonical data and store its output."""

    db: Database
    receipts: ReceiptRepository
    blobs: BlobStore
    history: HistorianQuery


def build_context(db: Database, blobs: BlobStore) -> ExportContext:
    return ExportContext(
        db=db,
        receipts=ReceiptRepository(db),
        blobs=blobs,
        history=HistorianQuery(db),
    )


async def fetch_receipts(
    ctx: ExportContext, user_id: str, params: Mapping[str, Any]
) -> tuple[Receipt, ...]:
    """Canonical rows in scope. Filtering is by value, never by a second stored copy."""
    rows = await ctx.receipts.list_for_user(user_id, limit=int(params.get("limit", 10000)))
    group_id = params.get("group_id")
    if group_id:
        rows = tuple(r for r in rows if r.group_id == group_id)
    return rows


def receipt_row(receipt: Receipt) -> list[Any]:
    return [
        receipt.receipt_id,
        receipt.transaction_date.date().isoformat() if receipt.transaction_date else "",
        receipt.vendor_name,
        receipt.currency,
        str(receipt.total_amount) if receipt.total_amount is not None else "",
        str(receipt.vat_amount) if receipt.vat_amount is not None else "",
        receipt.group_id or "",
    ]


def build_workbook(
    receipts: tuple[Receipt, ...],
    *,
    export_id: str,
    extra_columns: tuple[str, ...] = (),
    extra_values=None,
) -> bytes:
    """The raw-data sheet, the live-formula summary sheet, and the hidden reference cell.

    **The summary sheet is Excel formulas pulling from the raw sheet, not baked values.**
    That technique is kept deliberately: the print-friendly view stays correct when the raw
    sheet is edited, and it is rebuilt only when the *structure* changes rather than on every
    value edit. V2's file-open-tolerance machinery is **not** carried forward and must not
    be reintroduced — the server generates a file and the user downloads their own copy, so
    there is no file the server and a user's Excel are ever both holding open.

    The hidden `_resibo_meta` cell carries `export_id`, which is what makes Reimport's
    three-way diff possible at all. This writer and Reimport's parser are a matched pair.
    """
    try:
        from openpyxl import Workbook  # noqa: PLC0415
    except ImportError as exc:
        raise errors.ProviderUnavailable(f"openpyxl is not installed ({exc})") from exc

    workbook = Workbook()
    data = workbook.active
    data.title = "Receipts"
    columns = (*EXPORT_COLUMNS, *extra_columns)
    data.append(list(columns))
    for receipt in receipts:
        row = receipt_row(receipt)
        if extra_values is not None:
            row += list(extra_values(receipt))
        data.append(row)

    summary = workbook.create_sheet("Summary")
    summary["A1"] = "Receipts"
    summary["B1"] = f"=COUNTA(Receipts!A2:A{len(receipts) + 1})"
    summary["A2"] = "Total"
    summary["B2"] = f"=SUM(Receipts!E2:E{len(receipts) + 1})"
    summary["A3"] = "VAT"
    summary["B3"] = f"=SUM(Receipts!F2:F{len(receipts) + 1})"

    meta = workbook.create_sheet(SNAPSHOT_SHEET)
    meta[SNAPSHOT_CELL] = f"{SNAPSHOT_PREFIX}{export_id}"
    meta.sheet_state = "hidden"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def store_artifact(ctx: ExportContext, data: bytes) -> BlobRef:
    """Exports live in the same content-addressed blob store as everything else.

    An export is stored as `ORIGINAL` — it is not an archival image and there is no
    re-encoding step for it, so its `physical_hash` and `logical_id` happen to coincide.
    That coincidence is not something any caller may rely on: resolution still goes through
    `blob_locations` like every other blob.
    """
    result = await ctx.blobs.put(data, codec=StorageCodec.ORIGINAL)
    if not result.ok or result.blob_ref is None:
        raise errors.ExportError(result.error_detail or "blob write failed")
    return result.blob_ref


async def record_snapshot(ctx: ExportContext, export_id: str, user_id: str) -> datetime:
    """Register this export as a reimport baseline.

    Stores the Historian event id current at generation time. Reimport reconstructs the
    baseline *values* from the append-only trail as of `generated_at`, so this row is a
    pointer, never a second copy of the data that could drift from it.
    """
    event_id = await ctx.history.latest_event_id()
    generated_at = utcnow()

    def _write(conn) -> None:
        conn.execute(
            "INSERT OR REPLACE INTO export_snapshots (export_id, user_id,"
            " historian_event_id, generated_at) VALUES (?,?,?,?)",
            (export_id, user_id, event_id, generated_at.isoformat()),
        )

    await ctx.db.transaction(_write)
    return generated_at


def new_export_id() -> str:
    return uuid.uuid4().hex


__all__ = [
    "EXPORT_COLUMNS",
    "ExportContext",
    "build_context",
    "build_workbook",
    "fetch_receipts",
    "new_export_id",
    "receipt_row",
    "record_snapshot",
    "store_artifact",
]
