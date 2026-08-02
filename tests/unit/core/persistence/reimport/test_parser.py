"""The workbook parser: the missing-reference-cell hook, and the degradation path.

`openpyxl` is an optional dependency. The tests that need it skip when it is absent; the test
that asserts the *degradation* runs either way, because "reimport reports itself unavailable
rather than crashing" is the behaviour that has to hold on an install without it.
"""

from __future__ import annotations

import pytest

from core.persistence.reimport import errors
from core.persistence.reimport.parser import (
    DATA_SHEET,
    SNAPSHOT_CELL,
    SNAPSHOT_PREFIX,
    SNAPSHOT_SHEET,
    openpyxl_available,
    parse_workbook,
)

needs_openpyxl = pytest.mark.skipif(
    not openpyxl_available(), reason="openpyxl is an optional dependency"
)


def test_a_missing_file_is_an_error_code_not_a_raise(tmp_path):
    parsed = parse_workbook(tmp_path / "nothing-here.xlsx")
    assert not parsed.ok
    assert parsed.error_code in {errors.UNREADABLE_FILE, errors.PARSER_UNAVAILABLE}


def test_absent_openpyxl_degrades_rather_than_crashing(tmp_path, monkeypatch):
    """Simulates the dependency being missing on an install that never reimports."""
    import builtins

    real_import = builtins.__import__

    def _no_openpyxl(name, *args, **kwargs):
        if name == "openpyxl":
            raise ImportError("simulated missing dependency")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", _no_openpyxl)
    parsed = parse_workbook(tmp_path / "anything.xlsx")
    assert not parsed.ok
    assert parsed.error_code == errors.PARSER_UNAVAILABLE


@needs_openpyxl
def test_a_stripped_reference_cell_fails_cleanly(tmp_path):
    """A user deleting the row or column that held the hidden cell. Never a silent
    full-canonical overwrite."""
    from openpyxl import Workbook

    path = tmp_path / "no-reference.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = DATA_SHEET
    sheet.append(["receipt_id", "vendor_name"])
    sheet.append(["rc1", "Vendor A"])
    workbook.save(path)

    parsed = parse_workbook(path)
    assert not parsed.ok
    assert parsed.error_code == errors.MISSING_SNAPSHOT_REFERENCE
    assert parsed.rows == ()


@needs_openpyxl
def test_a_well_formed_workbook_round_trips(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "edited.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = DATA_SHEET
    sheet.append(["receipt_id", "vendor_name", "total_amount"])
    sheet.append(["rc1", "Vendor A", "100.00"])
    meta = workbook.create_sheet(SNAPSHOT_SHEET)
    meta[SNAPSHOT_CELL] = f"{SNAPSHOT_PREFIX}exp-1"
    workbook.save(path)

    parsed = parse_workbook(path)
    assert parsed.ok
    assert parsed.snapshot_export_id == "exp-1"
    assert parsed.rows[0].receipt_id == "rc1"
    assert parsed.rows[0].values["vendor_name"] == "Vendor A"
    assert "receipt_id" not in parsed.rows[0].values
